"""Parse BTG Pactual checking account statements (.xls/.xlsx)."""

import hashlib
from datetime import datetime
from typing import cast

import xlrd

# Spreadsheet cell values as produced by xlrd / openpyxl
CellValue = str | int | float | bool | datetime | None
Transaction = tuple[datetime, float, str]

# Spreadsheet column indices (0-based) as produced by BTG statements
DATE_COL = 1
TRANSACTION_TYPE_COL = 3
DESCRIPTION_COL = 6
AMOUNT_COL = 10

SALDO_DIARIO_DESCRIPTION = "Saldo Diário"
CONTA_LABEL = "Conta:"
SALDO_ATUAL_LABEL = "Saldo atual:"


def load_rows(path: str) -> list[list[CellValue]]:
    """Return the first sheet's rows as lists of cell values.

    Supports both legacy .xls (via xlrd) and .xlsx (via openpyxl).

    Some files are .xlsx content saved with a .xls extension, so we
    check the real format by the file header instead of the extension.
    """
    with open(path, "rb") as f:
        header = f.read(4)

    if header[:2] == b"PK":  # zip container -> xlsx
        from openpyxl import load_workbook

        # Pass a file object so openpyxl inspects content, not the extension.
        with open(path, "rb") as f:
            wb = load_workbook(f, read_only=True, data_only=True)
            ws = wb.active
            assert ws is not None  # statements always have an active sheet
            # Some generators write a bogus sheet dimension (e.g. A1);
            # force openpyxl to recompute it from the actual data.
            if (ws.max_row, ws.max_column) == (1, 1):
                ws.reset_dimensions()
            # Must read rows before the file handle is closed: read_only
            # mode streams lazily from the archive.
            rows = cast(
                list[list[CellValue]],
                [list(row) for row in ws.iter_rows(values_only=True)],
            )
            wb.close()
        return rows

    # Legacy BIFF .xls
    xls_wb = xlrd.open_workbook(path)
    sh = xls_wb.sheet_by_index(0)
    return cast(list[list[CellValue]], [sh.row_values(r) for r in range(sh.nrows)])


def norm_number(value: CellValue) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace("R$", "").replace(" ", "")
    if not s:
        return None
    # Handle "1.234,56" (pt-BR) or "1234.56"
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def norm_datetime(value: CellValue) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M"):
            try:
                return datetime.strptime(value, fmt)
            except ValueError:
                pass
    return None


# BTG Pactual bank code
BANK_ID = "208"
ORG = "BTG Pactual"


def esc(text: str) -> str:
    """Escape XML special characters."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def ofx_date(dt: datetime) -> str:
    # OFX date format: YYYYMMDDHHMMSS[gmt offset:tz name]
    return dt.strftime("%Y%m%d%H%M%S") + "[-3:BRT]"


def make_fitid(dt: datetime, amount: float, memo: str) -> str:
    raw = f"{dt.isoformat()}|{amount}|{memo}"
    return hashlib.md5(raw.encode("utf-8")).hexdigest()[:22]


def read_statement(path: str) -> tuple[str, float | None, list[Transaction]]:
    rows = load_rows(path)

    account_id = ""
    balance: float | None = None
    transactions: list[Transaction] = []

    for row in rows:
        cells = ["" if c is None else str(c).strip() for c in row]

        if CONTA_LABEL in cells:
            idx = cells.index(CONTA_LABEL)
            if idx + 1 < len(cells):
                account_id = cells[idx + 1]
        if SALDO_ATUAL_LABEL in cells:
            # The balance may be one or more columns after the label
            idx = cells.index(SALDO_ATUAL_LABEL)
            for v in row[idx + 1 :]:
                n = norm_number(v)
                if n is not None:
                    balance = n
                    break

        dt = norm_datetime(row[DATE_COL]) if len(row) > DATE_COL else None
        amount = norm_number(row[AMOUNT_COL]) if len(row) > AMOUNT_COL else None
        if dt is None or amount is None:
            continue

        transaction_type = (
            cells[TRANSACTION_TYPE_COL] if len(cells) > TRANSACTION_TYPE_COL else ""
        )
        description = cells[DESCRIPTION_COL] if len(cells) > DESCRIPTION_COL else ""

        # Skip daily-balance rows; they are not transactions
        if description == SALDO_DIARIO_DESCRIPTION:
            continue

        memo = " - ".join(p for p in (transaction_type, description) if p)
        transactions.append((dt, amount, memo))

    return account_id, balance, transactions


def trn_type(amount: float, memo: str) -> str:
    m = memo.lower()
    if "pix" in m:
        return "DEBIT" if amount < 0 else "CREDIT"
    if "salário" in m or "pagamento" in m:
        return "CREDIT" if amount >= 0 else "DEBIT"
    return "DEBIT" if amount < 0 else "CREDIT"
